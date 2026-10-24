"""Storage and exact anchor parsers for supported formats (SRS-005/006, AT-001/002).
Formats supported:
- text / markdown -> TEXT_LINES (1-based lines, char offsets)
- csv / tsv -> TABLE_CELLS (1-based row, col, header)
- docx -> DOCUMENT_BLOCK (block index, paragraph/table type)
- xlsx -> TABLE_CELLS (sheet name, cell coordinates, merged cell handling)
- pptx -> SLIDE_SHAPE (1-based slide, shape id, text)
- pdf -> PDF_PAGE (1-based page, text, fallback detection if image-only/needs OCR)
"""
from __future__ import annotations

import csv
import hashlib
import io
import os
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import docx
import openpyxl
import pptx
import pypdf

from .security import Detected, Rejected, detect

PARSER_VERSION = "1.0.0"


@dataclass(frozen=True)
class ParsedSpan:
    ordinal: int
    locator_type: str
    locator: dict[str, Any]
    text: str
    heading: str | None = None
    quality: str = "EXACT"  # EXACT | UNCERTAIN | OCR
    quality_note: str | None = None


@dataclass(frozen=True)
class ParseResult:
    kind: str
    parser: str
    parser_version: str
    status: str  # PARSED | PARTIAL | NEEDS_OCR | FAILED
    spans: list[ParsedSpan]
    diagnostics: list[str]


# ── Parser implementations ───────────────────────────────────────────────────

def _parse_text(data: bytes) -> ParseResult:
    text = data.decode("utf-8", errors="replace")
    lines = text.splitlines(keepends=True)
    spans: list[ParsedSpan] = []
    char_offset = 0
    current_heading = None

    for idx, line in enumerate(lines, start=1):
        stripped = line.strip()
        if not stripped:
            char_offset += len(line)
            continue
        if stripped.startswith("#"):
            current_heading = stripped.lstrip("#").strip()
        spans.append(
            ParsedSpan(
                ordinal=len(spans),
                locator_type="TEXT_LINES",
                locator={
                    "type": "TEXT_LINES",
                    "start_line": idx,
                    "end_line": idx,
                    "start_char": char_offset,
                    "end_char": char_offset + len(line),
                },
                text=stripped,
                heading=current_heading,
            )
        )
        char_offset += len(line)

    return ParseResult(
        kind="text",
        parser="dtp.text_parser",
        parser_version=PARSER_VERSION,
        status="PARSED" if spans else "PARTIAL",
        spans=spans,
        diagnostics=[] if spans else ["Document contained no non-empty lines"],
    )


def _parse_csv(data: bytes) -> ParseResult:
    text = data.decode("utf-8", errors="replace")
    dialect = "excel"
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample)
    except Exception:
        pass

    reader = csv.reader(io.StringIO(text), dialect)
    spans: list[ParsedSpan] = []
    header: list[str] = []

    for row_idx, row in enumerate(reader, start=1):
        if not row or not any(cell.strip() for cell in row):
            continue
        if row_idx == 1 and not header:
            header = [c.strip() for c in row]
            for col_idx, cell in enumerate(row, start=1):
                if cell.strip():
                    spans.append(
                        ParsedSpan(
                            ordinal=len(spans),
                            locator_type="TABLE_CELLS",
                            locator={
                                "type": "TABLE_CELLS",
                                "sheet": "default",
                                "row": row_idx,
                                "col": col_idx,
                                "is_header": True,
                            },
                            text=cell.strip(),
                            heading="Header",
                        )
                    )
            continue

        for col_idx, cell in enumerate(row, start=1):
            cell_val = cell.strip()
            if cell_val:
                col_name = header[col_idx - 1] if col_idx - 1 < len(header) else f"Col_{col_idx}"
                spans.append(
                    ParsedSpan(
                        ordinal=len(spans),
                        locator_type="TABLE_CELLS",
                        locator={
                            "type": "TABLE_CELLS",
                            "sheet": "default",
                            "row": row_idx,
                            "col": col_idx,
                            "col_name": col_name,
                        },
                        text=f"{col_name}: {cell_val}",
                        heading=f"Row {row_idx}",
                    )
                )

    return ParseResult(
        kind="csv",
        parser="dtp.csv_parser",
        parser_version=PARSER_VERSION,
        status="PARSED" if spans else "PARTIAL",
        spans=spans,
        diagnostics=[] if spans else ["CSV yielded no rows"],
    )


def _parse_docx(data: bytes) -> ParseResult:
    doc = docx.Document(io.BytesIO(data))
    spans: list[ParsedSpan] = []
    current_heading = None

    for p_idx, p in enumerate(doc.paragraphs, start=1):
        txt = p.text.strip()
        if not txt:
            continue
        if p.style and p.style.name and p.style.name.startswith("Heading"):
            current_heading = txt
        spans.append(
            ParsedSpan(
                ordinal=len(spans),
                locator_type="DOCUMENT_BLOCK",
                locator={
                    "type": "DOCUMENT_BLOCK",
                    "block_type": "paragraph",
                    "index": p_idx,
                    "style": p.style.name if p.style else None,
                },
                text=txt,
                heading=current_heading,
            )
        )

    for t_idx, table in enumerate(doc.tables, start=1):
        for r_idx, row in enumerate(table.rows, start=1):
            row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_texts:
                combined = " | ".join(dict.fromkeys(row_texts))  # de-duplicate merged cells
                spans.append(
                    ParsedSpan(
                        ordinal=len(spans),
                        locator_type="DOCUMENT_BLOCK",
                        locator={
                            "type": "DOCUMENT_BLOCK",
                            "block_type": "table_row",
                            "table_index": t_idx,
                            "row_index": r_idx,
                        },
                        text=combined,
                        heading=f"Table {t_idx}",
                    )
                )

    return ParseResult(
        kind="docx",
        parser="dtp.docx_parser",
        parser_version=PARSER_VERSION,
        status="PARSED" if spans else "PARTIAL",
        spans=spans,
        diagnostics=[],
    )


def _parse_xlsx(data: bytes) -> ParseResult:
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    spans: list[ParsedSpan] = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        merged_ranges = ws.merged_cells.ranges
        for r_idx, row in enumerate(ws.iter_rows(values_only=False), start=1):
            for c_idx, cell in enumerate(row, start=1):
                val = cell.value
                if val is not None and str(val).strip():
                    merged_info = None
                    for rng in merged_ranges:
                        if cell.coordinate in rng:
                            merged_info = str(rng)
                            break
                    spans.append(
                        ParsedSpan(
                            ordinal=len(spans),
                            locator_type="TABLE_CELLS",
                            locator={
                                "type": "TABLE_CELLS",
                                "sheet": sheet_name,
                                "row": r_idx,
                                "col": c_idx,
                                "coordinate": cell.coordinate,
                                "merged_range": merged_info,
                            },
                            text=str(val).strip(),
                            heading=f"{sheet_name}!{cell.coordinate}",
                        )
                    )

    return ParseResult(
        kind="xlsx",
        parser="dtp.xlsx_parser",
        parser_version=PARSER_VERSION,
        status="PARSED" if spans else "PARTIAL",
        spans=spans,
        diagnostics=[],
    )


def _parse_pptx(data: bytes) -> ParseResult:
    prs = pptx.Presentation(io.BytesIO(data))
    spans: list[ParsedSpan] = []

    for s_idx, slide in enumerate(prs.slides, start=1):
        slide_title = None
        if slide.shapes.title and slide.shapes.title.text.strip():
            slide_title = slide.shapes.title.text.strip()

        for sh_idx, shape in enumerate(slide.shapes, start=1):
            if shape.has_text_frame:
                txt = shape.text_frame.text.strip()
                if txt:
                    spans.append(
                        ParsedSpan(
                            ordinal=len(spans),
                            locator_type="SLIDE_SHAPE",
                            locator={
                                "type": "SLIDE_SHAPE",
                                "slide": s_idx,
                                "shape_id": shape.shape_id if hasattr(shape, "shape_id") else sh_idx,
                            },
                            text=txt,
                            heading=slide_title or f"Slide {s_idx}",
                        )
                    )

    return ParseResult(
        kind="pptx",
        parser="dtp.pptx_parser",
        parser_version=PARSER_VERSION,
        status="PARSED" if spans else "PARTIAL",
        spans=spans,
        diagnostics=[],
    )


def _parse_pdf(data: bytes) -> ParseResult:
    reader = pypdf.PdfReader(io.BytesIO(data))
    spans: list[ParsedSpan] = []
    total_pages = len(reader.pages)
    pages_with_text = 0

    for p_idx, page in enumerate(reader.pages, start=1):
        try:
            txt = page.extract_text() or ""
        except Exception:
            txt = ""
        stripped = txt.strip()
        if stripped:
            pages_with_text += 1
            spans.append(
                ParsedSpan(
                    ordinal=len(spans),
                    locator_type="PDF_PAGE",
                    locator={
                        "type": "PDF_PAGE",
                        "page": p_idx,
                        "total_pages": total_pages,
                    },
                    text=stripped,
                    heading=f"Page {p_idx}",
                )
            )

    diagnostics = []
    if total_pages > 0 and pages_with_text == 0:
        # Scanned PDF with no text layer -> requires OCR
        return ParseResult(
            kind="pdf",
            parser="dtp.pdf_parser",
            parser_version=PARSER_VERSION,
            status="NEEDS_OCR",
            spans=[],
            diagnostics=["PDF contains no extractable text layer (scanned/raster). Needs OCR adapter."],
        )
    elif pages_with_text < total_pages:
        diagnostics.append(f"{total_pages - pages_with_text} page(s) contained no text layer")

    return ParseResult(
        kind="pdf",
        parser="dtp.pdf_parser",
        parser_version=PARSER_VERSION,
        status="PARSED" if spans else "PARTIAL",
        spans=spans,
        diagnostics=diagnostics,
    )


def parse_bytes(data: bytes, declared_name: str) -> ParseResult:
    """Security detection gate followed by exact-anchor parser."""
    det = detect(data, declared_name)
    if det.kind in ("text", "csv"):
        if det.kind == "csv":
            return _parse_csv(data)
        return _parse_text(data)
    elif det.kind == "docx":
        return _parse_docx(data)
    elif det.kind == "xlsx":
        return _parse_xlsx(data)
    elif det.kind == "pptx":
        return _parse_pptx(data)
    elif det.kind == "pdf":
        return _parse_pdf(data)
    elif det.kind == "image":
        return ParseResult(
            kind="image",
            parser="dtp.image_intake",
            parser_version=PARSER_VERSION,
            status="NEEDS_OCR",
            spans=[],
            diagnostics=["Raw image requires OCR processing"],
        )
    raise Rejected("UNSUPPORTED_TYPE", f"Parser for {det.kind} not implemented")
